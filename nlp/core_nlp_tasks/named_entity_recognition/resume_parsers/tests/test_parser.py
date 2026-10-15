from io import BytesIO
import json
from types import SimpleNamespace

import fitz
from docx import Document
import pytest

from nlp.core_nlp_tasks.named_entity_recognition.resume_parsers.parser import (
    _validate_analysis,
    analyze_resume,
    extract_resume_text,
)


def valid_analysis():
    return {
        "parsed_data": {
            "contact_info": {
                "full_name": "Jordan Candidate",
                "email": "jordan@example.com",
                "phone": None,
                "location": None,
                "links": {"linkedin": None, "github": None, "portfolio": None},
            },
            "summary": None,
            "experience": [
                {
                    "company": "Example Co",
                    "title": "Data Analyst",
                    "duration": "2022-2025",
                    "highlights": ["Built SQL reporting dashboards."],
                }
            ],
            "education": [],
            "skills": {
                "technical": ["SQL"],
                "soft_skills": [],
                "tools_and_technologies": ["Power BI"],
                "languages": ["English"],
            },
            "certifications": [],
            "projects": [],
        },
        "extracted_keywords": {
            "primary_keywords": ["SQL", "Power BI"],
            "domain_expertise": ["Data analytics"],
        },
        "resume_quality_assessment": {
            "presentation_rating_1_to_10": 8,
            "career_level": "Mid-Level",
            "key_strengths": ["Reporting"],
            "weaknesses_or_gaps": ["No impact metrics stated"],
        },
        "job_fit_analysis": {
            "suitable_job_roles": ["Data Analyst"],
            "qualification_status": "Moderately Qualified",
            "suitability_reasoning": "SQL experience is relevant; the resume gives limited evidence of impact.",
        },
    }


def test_extract_resume_text_supports_utf8_text():
    assert extract_resume_text("resume.txt", b"\xef\xbb\xbfData analyst\nSQL") == "Data analyst\nSQL"


def test_extract_resume_text_supports_pdf():
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "Jordan Candidate - Data Analyst with experience building SQL reporting dashboards.",
    )
    data = document.tobytes()
    document.close()

    assert "SQL reporting dashboards" in extract_resume_text("resume.pdf", data)


def test_extract_resume_text_supports_docx_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("Jordan Candidate")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Experience"
    table.cell(0, 1).text = "Data Analyst"
    output = BytesIO()
    document.save(output)

    text = extract_resume_text("resume.docx", output.getvalue())
    assert "Jordan Candidate" in text
    assert "Experience | Data Analyst" in text


def test_extract_resume_text_rejects_unsupported_file_type():
    with pytest.raises(ValueError, match="Unsupported resume format"):
        extract_resume_text("resume.rtf", b"resume")


def test_analysis_validation_rejects_invalid_quality_rating():
    result = valid_analysis()
    result["resume_quality_assessment"]["presentation_rating_1_to_10"] = 11

    with pytest.raises(ValueError, match="integer from 1 to 10"):
        _validate_analysis(result)


def test_analyze_resume_requests_json_and_validates_response():
    expected = valid_analysis()

    class FakeClient:
        def __init__(self):
            self.chat = SimpleNamespace(completions=SimpleNamespace(create=self.create))
            self.request = None

        def create(self, **kwargs):
            self.request = kwargs
            content = json.dumps(expected)
            return SimpleNamespace(
                choices=[SimpleNamespace(message=SimpleNamespace(content=content))]
            )

    client = FakeClient()
    result = analyze_resume(
        "Jordan Candidate has SQL and Power BI experience.",
        "Senior Data Analyst",
        "Must have SQL and analytics experience.",
        lambda: (client, "test-model"),
    )

    assert result == expected
    assert client.request["response_format"] == {"type": "json_object"}
    assert "Senior Data Analyst" in client.request["messages"][1]["content"]
    assert "analytics experience" in client.request["messages"][1]["content"]
